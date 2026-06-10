import os
import pandas as pd
import datetime
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font

from config import OUTPUT_FILE


class ResultHandler:

    def __init__(self):
        self.data = []

    @staticmethod
    def _normalize_price(value):
        if value is None:
            return ""

        text = str(value).strip()
        if not text:
            return ""

        text = text.replace("€", "")
        text = text.replace("EUR", "")
        text = text.replace("euro", "")
        text = text.replace(" ", "")

        # Handle Dutch style like "1.149,-"
        if text.endswith(",-"):
            text = text[:-2]
            text = text.replace(".", "")
            text += ".00"

        # Normalize decimal separators
        if "," in text and "." in text:
            if text.rfind(",") > text.rfind("."):
                text = text.replace(".", "").replace(",", ".")
            else:
                text = text.replace(",", "")
        elif "," in text:
            parts = text.split(",")
            if len(parts) == 2 and len(parts[1]) in (1, 2):
                text = text.replace(",", ".")
            else:
                text = text.replace(",", "")
        elif "." in text:
            parts = text.split(".")
            if len(parts) == 2 and len(parts[1]) in (1, 2):
                pass
            else:
                text = text.replace(".", "")

        # Keep only digits, minus sign and decimal point
        text = "".join(ch for ch in text if ch.isdigit() or ch in "-.")

        if not text or text in {"-", "."}:
            return ""

        # Force plain numeric format
        if "." not in text:
            text += ".00"
        else:
            parts = text.split(".")
            if len(parts) == 2 and len(parts[1]) == 1:
                text += "0"
            elif len(parts) == 2 and len(parts[1]) == 0:
                text += "00"
            elif len(parts) == 2 and len(parts[1]) > 2:
                text = parts[0] + "." + parts[1][:2]

        return text

    # ===========================
    # ADD RESULT
    # ===========================
    def add(self, result):
        if result is None:
            return

        self.data.append(result)

    # ===========================
    # SAVE (PIVOT FORMAT)
    # ===========================
    def save(self):
        if not self.data:
            print("No data to save.")
            return

        df = pd.DataFrame(self.data)

        if "price" in df.columns:
            df["price"] = df["price"].apply(self._normalize_price)

        if "timestamp" in df.columns:
            df = df.drop(columns=["timestamp"])

        # ✅ pivot: site → columns
        pivot_df = df.pivot_table(
            index="model",
            columns="site",
            values="price",
            aggfunc="first"
        )

        # ✅ flatten columns
        pivot_df.columns.name = None
        pivot_df.reset_index(inplace=True)

        # ✅ column order
        cols = ["model"] + [c for c in pivot_df.columns if c != "model"]
        final_df = pivot_df[cols]

        # ✅ filename
        base_filename = OUTPUT_FILE.replace(".xlsx", "")
        date_str = datetime.datetime.now().strftime("%Y-%m-%d")

        version = 1
        while True:
            filename = f"{base_filename}_{date_str}_v{version}.xlsx"
            if not os.path.exists(filename):
                break
            version += 1



        final_df.to_excel(filename, index=False)

        workbook = load_workbook(filename)
        sheet = workbook.active

        # Highlight header row
        header_fill = PatternFill("solid", fgColor="D9EAF7")
        header_font = Font(bold=True, color="000000")
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font

        # Set column width for all columns
        for i, column in enumerate(sheet.columns, start=1):
            col_letter = column[0].column_letter

            if i == 1:
                sheet.column_dimensions[col_letter].width = 25
            else:
                sheet.column_dimensions[col_letter].width = 13


        workbook.save(filename)

        print(f"\nSaved results: {filename}")